import os
import shutil
import json
from io import BytesIO
from datetime import datetime, timezone
from fastapi import File, UploadFile
from typing import Any
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.db import SessionLocal, get_session
from app.deps import Principal, require_admin
from app.engine import (
    EngineError,
    admin_question_matrix,
    audit,
    close_question,
    get_round,
    leaderboard,
    override_grade,
    pause_question,
    published_leaderboard,
    reorder_questions,
    resume_question,
    restart_round,
    show_topic,
    start_question,
)
from app.models import AuditEvent, AuthSession, PublishedLeaderboard, Question, RoundState, SwapEvent, Team, TeamQuestion, User
from app.redisutil import get_controller, release_controller
from app.security import hash_password
from app.ws import ADMIN, PUBLIC, hub, notify_after_commit, team_channel

router = APIRouter(prefix="/api/v1/admin", tags=["admin"])


async def get_export_session():
    """Provide an isolated, non-mutating snapshot transaction for exports."""
    async with SessionLocal() as session:
        bind = session.get_bind()
        if bind.dialect.name == "postgresql":
            await session.execute(
                text("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
            )
        try:
            yield session
        finally:
            # The export transaction is deliberately never committed.
            await session.rollback()


def _json_default(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, UUID):
        return str(value)
    raise TypeError(f"Object of type {type(value).__name__} is not JSON serializable")


def _excel_value(value: Any) -> Any:
    """Convert database values into safe, readable spreadsheet cells."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, (dict, list)):
        value = json.dumps(value, default=_json_default, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _add_excel_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([_excel_value(value) for value in row])
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 60)
    return sheet


def _session_export_status(row: AuthSession, exported_at: datetime) -> str:
    if row.revoked_at:
        return "REVOKED"
    expires_at = row.expires_at
    if expires_at.tzinfo is None:
        expires_at = expires_at.replace(tzinfo=timezone.utc)
    return "EXPIRED" if expires_at <= exported_at else "VALID"


class TeamIn(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    username: str = Field(min_length=1, max_length=64)
    password: str = Field(min_length=3, max_length=128)
    score: int | None = None


class ScoreIn(BaseModel):
    score: int


class QuestionIn(BaseModel):
    index: int = Field(ge=1, le=99)
    topic: str
    prompt: str
    qtype: str = Field(pattern="^(mcq_single|mcq_multi)$")
    options: Any | None = None
    answer_key: dict[str, Any]
    hint: str | None = None
    media_url: str | None = None


class QuestionOrderIn(BaseModel):
    question_ids: list[UUID]


class OverrideIn(BaseModel):
    team_id: UUID
    question_id: UUID
    correct: bool


class RestartIn(BaseModel):
    confirmation: str
    expected_round_version: int = Field(ge=0)


def _err(e: EngineError) -> HTTPException:
    return HTTPException(e.status, {"code": e.code, "message": e.message})


async def _fanout(
    session: AsyncSession,
    extra_admin: dict | None = None,
) -> None:
    from app.engine import get_round as gr
    from app.models import Question as Q

    rnd = await gr(session)
    live_board = await leaderboard(session)
    frozen_board = await published_leaderboard(session)

    q = await session.get(Q, rnd.current_question_id) if rnd.current_question_id else None
    
    # PUBLIC gets the FROZEN leaderboard
    public = {
        "type": "state",
        "phase": rnd.phase,
        "server_now": None,
        "topic_shown_at": rnd.topic_shown_at.isoformat() if rnd.topic_shown_at else None,
        "live_started_at": rnd.live_started_at.isoformat() if rnd.live_started_at else None,
        "global_deadline_at": rnd.global_deadline_at.isoformat() if rnd.global_deadline_at else None,
        "paused_at": rnd.paused_at.isoformat() if rnd.paused_at else None,
        "question": (
            {
                "id": str(q.id),
                "index": q.index,
                "topic": q.topic,
                "type": q.qtype,
                "prompt": q.prompt if rnd.phase in ("live", "paused", "closed") else None,
                "media_url": q.media_url if rnd.phase in ("live", "paused", "closed") else None,
                "options": q.options if rnd.phase in ("live", "paused", "closed") else None,
            }
            if q
            else None
        ),
        "leaderboard": frozen_board, 
    }
    from datetime import datetime, timezone

    public["server_now"] = datetime.now(timezone.utc).isoformat()
    await notify_after_commit(
        hub.broadcast(PUBLIC, public),
        operation="admin.fanout.public",
        context={"phase": rnd.phase},
    )
    
    # ADMIN payload gets the LIVE leaderboard so you can monitor point spending in real-time
    admin_payload = dict(public)
    admin_payload["type"] = "admin_state"
    admin_payload["leaderboard"] = live_board 
    if extra_admin:
        admin_payload.update(extra_admin)
        
    await notify_after_commit(
        hub.broadcast(ADMIN, admin_payload),
        operation="admin.fanout.admin",
        context={"phase": rnd.phase},
    )
    
    if q:
        matrix = await admin_question_matrix(session, q.id)
        await notify_after_commit(
            hub.broadcast(ADMIN, {"type": "matrix", "matrix": matrix}),
            operation="admin.fanout.matrix",
            context={"question_id": str(q.id)},
        )
        for t in live_board:
            # Tell teams to update, but only give them the frozen board
            await notify_after_commit(
                hub.broadcast(team_channel(t["id"]), {"type": "state_tick", "phase": rnd.phase, "leaderboard": frozen_board}),
                operation="admin.fanout.team",
                context={"team_id": t["id"], "phase": rnd.phase},
            )

@router.post("/teams")
async def create_team(
    body: TeamIn,
    session: AsyncSession = Depends(get_session),
    _: Principal = Depends(require_admin),
):
    if (await session.execute(select(User).where(User.username == body.username))).scalar_one_or_none():
        raise HTTPException(409, "Username taken")
    if (await session.execute(select(Team).where(Team.name == body.name))).scalar_one_or_none():
        raise HTTPException(409, "Team name taken")
    settings = get_settings()
    team = Team(name=body.name, score=body.score if body.score is not None else settings.initial_team_score)
    session.add(team)
    await session.flush()
    user = User(
        username=body.username,
        password_hash=hash_password(body.password),
        role="team",
        team_id=team.id,
    )
    session.add(user)
    await session.commit()
    return {"team_id": str(team.id), "username": body.username, "name": team.name, "score": team.score}


@router.get("/teams")
async def list_teams(session: AsyncSession = Depends(get_session), _: Principal = Depends(require_admin)):
    teams = (await session.execute(select(Team).order_by(Team.name))).scalars().all()
    return [{"id": str(t.id), "name": t.name, "score": t.score, "swap_available": t.swap_available} for t in teams]


def _utc(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def _session_activity(row: AuthSession, now: datetime) -> tuple[str, float, bool]:
    now = _utc(now)
    age_seconds = max(0.0, (now - _utc(row.last_seen_at)).total_seconds())
    expired = _utc(row.expires_at) <= now
    structurally_valid = row.team_id is not None and row.slot is not None
    valid = row.revoked_at is None and not expired and structurally_valid

    if row.revoked_at is not None:
        status = "revoked"
    elif expired:
        status = "expired"
    elif age_seconds <= 90:
        status = "active"
    elif age_seconds <= 180:
        status = "stale"
    else:
        status = "inactive"
    return status, round(age_seconds, 3), valid


@router.get("/teams/{team_id}/sessions")
async def list_team_sessions(
    team_id: UUID,
    session: AsyncSession = Depends(get_session),
    _: Principal = Depends(require_admin),
):
    if await session.get(Team, team_id) is None:
        raise HTTPException(404, "Team not found")
    rows = (
        await session.execute(
            select(AuthSession)
            .where(AuthSession.team_id == team_id)
            .order_by(AuthSession.created_at.desc())
        )
    ).scalars().all()
    now = datetime.now(timezone.utc)
    controller_available = True
    try:
        controller_sid = await get_controller(str(team_id))
    except Exception:
        controller_available = False
        controller_sid = None
    activity = {row.id: _session_activity(row, now) for row in rows}
    payload = [
        {
            "session_id": str(row.id),
            "device_label": row.device_label,
            "user_agent": row.user_agent,
            "login_time": row.created_at.isoformat(),
            "last_seen": row.last_seen_at.isoformat(),
            "expires_at": row.expires_at.isoformat(),
            "revoked_at": row.revoked_at.isoformat() if row.revoked_at else None,
            "status": activity[row.id][0],
            "last_seen_age_sec": activity[row.id][1],
            "valid": activity[row.id][2],
            # Compatibility: historically `active` meant authentication-valid.
            "active": activity[row.id][2],
            "is_controller": (
                controller_sid == str(row.id) if controller_available else None
            ),
        }
        for row in rows
    ]
    valid_session_count = sum(1 for status in activity.values() if status[2])
    recent_session_count = sum(1 for status in activity.values() if status[0] == "active")
    return {
        "team_id": str(team_id),
        # Compatibility for existing clients; this remains the valid-session count.
        "active_count": valid_session_count,
        "valid_session_count": valid_session_count,
        "recent_session_count": recent_session_count,
        "controller_available": controller_available,
        "controller_session_id": controller_sid,
        "sessions": payload,
    }


@router.post("/teams/{team_id}/sessions/revoke-all")
async def revoke_all_team_sessions(
    team_id: UUID,
    session: AsyncSession = Depends(get_session),
    _: Principal = Depends(require_admin),
):
    if await session.get(Team, team_id) is None:
        raise HTTPException(404, "Team not found")
    rows = (
        await session.execute(
            select(AuthSession)
            .where(
                AuthSession.team_id == team_id,
                AuthSession.revoked_at.is_(None),
            )
            .with_for_update()
        )
    ).scalars().all()
    now = datetime.now(timezone.utc)
    for row in rows:
        row.revoked_at = now
        row.slot = None
    await session.commit()
    return {"ok": True, "revoked_count": len(rows)}


@router.post("/teams/{team_id}/sessions/{session_id}/revoke")
async def revoke_team_session(
    team_id: UUID,
    session_id: UUID,
    session: AsyncSession = Depends(get_session),
    _: Principal = Depends(require_admin),
):
    row = (
        await session.execute(
            select(AuthSession)
            .where(AuthSession.id == session_id, AuthSession.team_id == team_id)
            .with_for_update()
        )
    ).scalar_one_or_none()
    if row is None:
        raise HTTPException(404, "Session not found")
    if row.revoked_at is None:
        row.revoked_at = datetime.now(timezone.utc)
        row.slot = None
        await session.commit()
    return {"ok": True, "session_id": str(session_id)}


@router.patch("/teams/{team_id}/score")
async def set_score(
    team_id: UUID,
    body: ScoreIn,
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    rnd = await get_round(session, for_update=True)
    team = (
        await session.execute(
            select(Team).where(Team.id == team_id).with_for_update()
        )
    ).scalar_one_or_none()
    if team is None:
        raise HTTPException(404, "Team not found")
    old_score = team.score
    team.score = body.score
    team.version += 1
    await audit(
        session,
        actor=admin.username,
        kind="score_set",
        team_id=team.id,
        question_id=rnd.current_question_id,
        payload={
            "old_score": old_score,
            "new_score": body.score,
            "delta": body.score - old_score,
            "round_version": rnd.version,
            "phase": rnd.phase,
            "current_question_id": (
                str(rnd.current_question_id) if rnd.current_question_id else None
            ),
        },
    )
    await session.commit()
    await _fanout(session)
    return {"id": str(team.id), "score": team.score}


@router.post("/questions")
async def create_question(
    body: QuestionIn,
    session: AsyncSession = Depends(get_session),
    _: Principal = Depends(require_admin),
):
    if (await session.execute(select(Question).where(Question.index == body.index))).scalar_one_or_none():
        raise HTTPException(409, "Question index already exists")
    q = Question(
        index=body.index,
        topic=body.topic,
        prompt=body.prompt,
        qtype=body.qtype,
        options=body.options,
        answer_key=body.answer_key,
        hint=body.hint,
        media_url = body.media_url,
    )
    session.add(q)
    await session.commit()
    await session.refresh(q)
    return {"id": str(q.id), "index": q.index}


@router.post("/questions/reorder")
async def reorder_question_sequence(
    body: QuestionOrderIn,
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    try:
        await reorder_questions(session, body.question_ids, actor=admin.username)
        await session.commit()
    except EngineError as e:
        await session.rollback()
        raise _err(e) from None
    return {"ok": True}


@router.put("/questions/{question_id}")
async def update_question(
    question_id: UUID,
    body: QuestionIn,
    session: AsyncSession = Depends(get_session),
    _: Principal = Depends(require_admin),
):
    q = await session.get(Question, question_id)
    if q is None:
        raise HTTPException(404, "Not found")
    if body.index != q.index:
        dup = (await session.execute(select(Question).where(Question.index == body.index, Question.id != question_id))).scalar_one_or_none()
        if dup is not None:
            raise HTTPException(409, "Question index already exists")
    q.index = body.index
    q.topic = body.topic
    q.prompt = body.prompt
    q.qtype = body.qtype
    q.options = body.options
    q.answer_key = body.answer_key
    q.hint = body.hint
    q.media_url = body.media_url
    await session.commit()
    return {"id": str(q.id)}


@router.get("/questions")
async def list_questions(session: AsyncSession = Depends(get_session), _: Principal = Depends(require_admin)):
    qs = (await session.execute(select(Question).order_by(Question.index))).scalars().all()
    return [
        {
            "id": str(q.id),
            "index": q.index,
            "topic": q.topic,
            "prompt": q.prompt,
            "type": q.qtype,
            "options": q.options,
            "answer_key": q.answer_key,
            "hint": q.hint,
            "media_url": q.media_url,
        }
        for q in qs
    ]


@router.post("/round/show-topic/{question_id}")
async def api_show_topic(
    question_id: UUID,
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    try:
        await show_topic(session, question_id, actor=admin.username)
        await session.commit()
    except EngineError as e:
        await session.rollback()
        raise _err(e) from None
    await _fanout(session)
    return {"ok": True}


@router.post("/round/start")
async def api_start(
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    try:
        await start_question(session, actor=admin.username)
        await session.commit()
    except EngineError as e:
        await session.rollback()
        raise _err(e) from None
    await _fanout(session)
    return {"ok": True}


@router.post("/round/pause")
async def api_pause(
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    try:
        rnd, changed = await pause_question(session, actor=admin.username)
        if changed:
            await session.commit()
    except EngineError as e:
        await session.rollback()
        raise _err(e) from None
    if changed:
        await _fanout(session)
    return {
        "ok": True,
        "paused": True,
        "already_paused": not changed,
        "paused_at": rnd.paused_at.isoformat() if rnd.paused_at else None,
        "version": rnd.version,
    }


@router.post("/round/resume")
async def api_resume(
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    try:
        rnd, changed = await resume_question(session, actor=admin.username)
        if changed:
            await session.commit()
    except EngineError as e:
        await session.rollback()
        raise _err(e) from None
    if changed:
        await _fanout(session)
    return {
        "ok": True,
        "paused": False,
        "already_resumed": not changed,
        "version": rnd.version,
    }


@router.post("/round/close")
async def api_close(
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    try:
        rnd = await close_question(session, actor=admin.username)
        await session.commit()
    except EngineError as e:
        await session.rollback()
        raise _err(e) from None

    extra = {}
    if rnd.current_question_id:
        extra["matrix"] = await admin_question_matrix(session, rnd.current_question_id)
        extra["swap_notices"] = extra["matrix"]["swaps"]
        for s in extra["matrix"]["swaps"]:
            await notify_after_commit(
                hub.broadcast(
                    team_channel(s["target"]),
                    {"type": "swap_notice", "source": s["source"], "question_id": str(rnd.current_question_id)},
                ),
                operation="admin.close.swap_target",
                context={"target_team_id": s["target"], "question_id": str(rnd.current_question_id)},
            )
            await notify_after_commit(
                hub.broadcast(
                    team_channel(s["source"]),
                    {"type": "swap_applied", "target": s["target"], "question_id": str(rnd.current_question_id)},
                ),
                operation="admin.close.swap_source",
                context={"source_team_id": s["source"], "question_id": str(rnd.current_question_id)},
            )
    await _fanout(session, extra_admin=extra)
    return {"ok": True}


@router.post("/round/restart")
async def api_restart(
    body: RestartIn,
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    if body.confirmation != "RESTART ROUND":
        raise HTTPException(
            400,
            {"code": "invalid_confirmation", "message": "Confirmation phrase does not match"},
        )
    try:
        await restart_round(
            session,
            actor=admin.username,
            expected_round_version=body.expected_round_version,
        )
        await session.commit()
    except EngineError as e:
        await session.rollback()
        raise _err(e) from None

    from app.redisutil import clear_frozen_leaderboard

    await notify_after_commit(
        clear_frozen_leaderboard(),
        operation="admin.restart.clear_frozen_leaderboard",
        context={"round_version": body.expected_round_version},
    )
    reset_board = await leaderboard(session)
    for team in reset_board:
        await notify_after_commit(
            release_controller(team["id"], force=True),
            operation="admin.restart.release_controller",
            context={"team_id": team["id"]},
        )
    await _fanout(session)
    return {"ok": True}


@router.get("/round")
async def api_round(session: AsyncSession = Depends(get_session), _: Principal = Depends(require_admin)):
    rnd = await get_round(session)
    q = await session.get(Question, rnd.current_question_id) if rnd.current_question_id else None
    board = await leaderboard(session)
    matrix = None
    if rnd.current_question_id:
        matrix = await admin_question_matrix(session, rnd.current_question_id)
    from datetime import datetime, timezone, timedelta

    global_deadline_ms = None

    # everything converted to milliseconds
    if rnd.phase == "topic" and rnd.topic_shown_at:
        dt = rnd.topic_shown_at.replace(tzinfo=timezone.utc) if not rnd.topic_shown_at.tzinfo else rnd.topic_shown_at
        global_deadline_ms = int(dt.timestamp() * 1000) + 45000

    elif rnd.phase in ("live", "paused") and rnd.live_started_at:
        dt = rnd.live_started_at.replace(tzinfo=timezone.utc) if not rnd.live_started_at.tzinfo else rnd.live_started_at
        global_deadline_ms = int(dt.timestamp() * 1000) + 150000
        if rnd.global_deadline_at:
            g_dt = rnd.global_deadline_at.replace(tzinfo=timezone.utc) if not rnd.global_deadline_at.tzinfo else rnd.global_deadline_at
            global_deadline_ms = int(g_dt.timestamp() * 1000)

    return {
        "phase": rnd.phase,
        "version": rnd.version,
        "server_now": datetime.now(timezone.utc).isoformat(),
        "topic_shown_at": rnd.topic_shown_at.isoformat() if rnd.topic_shown_at else None,
        "live_started_at": rnd.live_started_at.isoformat() if rnd.live_started_at else None,
        "global_deadline_at": rnd.global_deadline_at.isoformat() if rnd.global_deadline_at else None,
        "paused_at": rnd.paused_at.isoformat() if rnd.paused_at else None,
        "global_deadline_ms": global_deadline_ms,
        "current_question_id": str(rnd.current_question_id) if rnd.current_question_id else None,
        "question": (
            {
                "id": str(q.id),
                "index": q.index,
                "topic": q.topic,
                "type": q.qtype,
                "prompt": q.prompt if rnd.phase in ("live", "paused", "closed") else None,
                "media_url": q.media_url if rnd.phase in ("live", "paused", "closed") else None,
                "options": q.options if rnd.phase in ("live", "paused", "closed") else None,
            }
            if q
            else None
        ),
        "leaderboard": board,
        "matrix": matrix,
        "reveal_cards": rnd.reveal_cards,
    }

@router.post("/grade/override")
async def api_override(
    body: OverrideIn,
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    try:
        await override_grade(session, body.team_id, body.question_id, body.correct, actor=admin.username)
        await session.commit()
    except EngineError as e:
        await session.rollback()
        raise _err(e) from None
    await _fanout(session)
    return {"ok": True}


@router.post("/teams/{team_id}/release-controller")
async def api_release(
    team_id: UUID,
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    await release_controller(str(team_id), force=True)
    await hub.broadcast(team_channel(str(team_id)), {"type": "controller_released"})
    return {"ok": True}


@router.get("/teams/{team_id}/audit")
async def api_team_audit(
    team_id: UUID,
    session: AsyncSession = Depends(get_session),
    _: Principal = Depends(require_admin),
):
    if await session.get(Team, team_id) is None:
        raise HTTPException(404, "Team not found")
    rows = (
        await session.execute(
            select(AuditEvent)
            .where(AuditEvent.team_id == team_id)
            .order_by(AuditEvent.at.asc())
        )
    ).scalars().all()
    return [
        {
            "at": r.at.isoformat() if r.at else None,
            "actor": r.actor,
            "kind": r.kind,
            "team_id": str(r.team_id) if r.team_id else None,
            "question_id": str(r.question_id) if r.question_id else None,
            "payload": r.payload,
        }
        for r in rows
    ]


@router.get("/audit")
async def api_audit(session: AsyncSession = Depends(get_session), _: Principal = Depends(require_admin)):
    rows = (
        await session.execute(select(AuditEvent).order_by(AuditEvent.at.desc()).limit(500))
    ).scalars().all()
    return [
        {
            "at": r.at.isoformat() if r.at else None,
            "actor": r.actor,
            "kind": r.kind,
            "team_id": str(r.team_id) if r.team_id else None,
            "question_id": str(r.question_id) if r.question_id else None,
            "payload": r.payload,
        }
        for r in rows
    ]


@router.get("/export/event")
async def export_event(
    _: Principal = Depends(require_admin),
    export_session: AsyncSession = Depends(get_export_session),
):
    """Download a complete, read-only recovery snapshot of the event."""
    exported_at = datetime.now(timezone.utc)

    rounds = (
        await export_session.execute(select(RoundState).order_by(RoundState.id))
    ).scalars().all()
    teams = (
        await export_session.execute(select(Team).order_by(Team.name, Team.id))
    ).scalars().all()
    users = (
        await export_session.execute(select(User).order_by(User.username, User.id))
    ).scalars().all()
    questions = (
        await export_session.execute(select(Question).order_by(Question.index, Question.id))
    ).scalars().all()
    team_questions = (
        await export_session.execute(
            select(TeamQuestion).order_by(TeamQuestion.question_id, TeamQuestion.team_id)
        )
    ).scalars().all()
    swaps = (
        await export_session.execute(
            select(SwapEvent).order_by(SwapEvent.created_at, SwapEvent.id)
        )
    ).scalars().all()
    audit_events = (
        await export_session.execute(
            select(AuditEvent).order_by(AuditEvent.at, AuditEvent.id)
        )
    ).scalars().all()
    auth_sessions = (
        await export_session.execute(
            select(AuthSession).order_by(AuthSession.created_at, AuthSession.id)
        )
    ).scalars().all()
    durable_published_leaderboard = await export_session.get(PublishedLeaderboard, 1)

    current_round = next((row for row in rounds if row.id == 1), None)
    current_question = None
    if current_round and current_round.current_question_id:
        question = next(
            (row for row in questions if row.id == current_round.current_question_id),
            None,
        )
        current_question = {
            "id": str(current_round.current_question_id),
            "index": question.index if question else None,
            "topic": question.topic if question else None,
        }

    postgres = {
        "round_states": [
            {
                "id": row.id,
                "phase": row.phase,
                "current_question_id": str(row.current_question_id) if row.current_question_id else None,
                "topic_shown_at": row.topic_shown_at,
                "live_started_at": row.live_started_at,
                "global_deadline_at": row.global_deadline_at,
                "paused_at": row.paused_at,
                "closed_at": row.closed_at,
                "reveal_cards": row.reveal_cards,
                "version": row.version,
            }
            for row in rounds
        ],
        "teams": [
            {
                "id": str(row.id),
                "name": row.name,
                "score": row.score,
                "swap_available": row.swap_available,
                "version": row.version,
                "created_at": row.created_at,
            }
            for row in teams
        ],
        "users": [
            {
                "id": str(row.id),
                "username": row.username,
                "role": row.role,
                "team_id": str(row.team_id) if row.team_id else None,
                "is_active": row.is_active,
                "created_at": row.created_at,
            }
            for row in users
        ],
        "questions": [
            {
                "id": str(row.id),
                "index": row.index,
                "topic": row.topic,
                "prompt": row.prompt,
                "qtype": row.qtype,
                "media_url": row.media_url,
                "options": row.options,
                "answer_key": row.answer_key,
                "hint": row.hint,
                "created_at": row.created_at,
            }
            for row in questions
        ],
        "team_questions": [
            {
                "id": str(row.id),
                "team_id": str(row.team_id),
                "question_id": str(row.question_id),
                "bounty": row.bounty,
                "freeze": row.freeze,
                "peek": row.peek,
                "swap_used": row.swap_used,
                "freeze_elapsed_ms": row.freeze_elapsed_ms,
                "locked": row.locked,
                "locked_elapsed_ms": row.locked_elapsed_ms,
                "deadline_at": row.deadline_at,
                "answer": row.answer,
                "kingslayer_eligible": row.kingslayer_eligible,
                "points_delta": row.points_delta,
                "graded": row.graded,
                "is_correct": row.is_correct,
                "admin_override": row.admin_override,
            }
            for row in team_questions
        ],
        "swap_events": [
            {
                "id": str(row.id),
                "question_id": str(row.question_id),
                "source_team_id": str(row.source_team_id),
                "target_team_id": str(row.target_team_id),
                "created_at": row.created_at,
                "applied": row.applied,
            }
            for row in swaps
        ],
        "audit_events": [
            {
                "id": str(row.id),
                "at": row.at,
                "actor": row.actor,
                "team_id": str(row.team_id) if row.team_id else None,
                "question_id": str(row.question_id) if row.question_id else None,
                "kind": row.kind,
                "payload": row.payload,
            }
            for row in audit_events
        ],
        "auth_sessions": [
            {
                "id": str(row.id),
                "user_id": str(row.user_id),
                "team_id": str(row.team_id) if row.team_id else None,
                "slot": row.slot,
                "created_at": row.created_at,
                "last_seen_at": row.last_seen_at,
                "expires_at": row.expires_at,
                "revoked_at": row.revoked_at,
                "device_label": row.device_label,
                "user_agent": row.user_agent,
            }
            for row in auth_sessions
        ],
        "published_leaderboard": (
            {
                "id": durable_published_leaderboard.id,
                "published_at": durable_published_leaderboard.published_at,
                "published_by": durable_published_leaderboard.published_by,
                "round_version": durable_published_leaderboard.round_version,
                "phase": durable_published_leaderboard.phase,
                "question_id": (
                    str(durable_published_leaderboard.question_id)
                    if durable_published_leaderboard.question_id
                    else None
                ),
                "question_index": durable_published_leaderboard.question_index,
                "rows": durable_published_leaderboard.rows,
                "card_usage_rows": durable_published_leaderboard.card_usage_rows,
                "published_solution": durable_published_leaderboard.published_solution,
            }
            if durable_published_leaderboard
            else None
        ),
    }

    try:
        from app.redisutil import get_frozen_leaderboard

        controllers = {
            str(team.id): await get_controller(str(team.id)) for team in teams
        }
        redis_state = {
            "available": True,
            "controllers": controllers,
            "published_leaderboard": await get_frozen_leaderboard(),
            "error": None,
        }
    except Exception:
        redis_state = {
            "available": False,
            "controllers": None,
            "published_leaderboard": None,
            "error": "unavailable",
        }

    payload = {
        "schema_version": 1,
        "exported_at": exported_at,
        "application": "levelup",
        "summary": {
            "current_phase": current_round.phase if current_round else None,
            "current_round_version": current_round.version if current_round else None,
            "current_question": current_question,
        },
        "postgres": postgres,
        "redis": redis_state,
    }
    body = json.dumps(payload, default=_json_default, ensure_ascii=False, indent=2)
    filename = f"levelup-event-{exported_at.strftime('%Y-%m-%dT%H-%M-%SZ')}.json"
    return Response(
        content=body,
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export/event.xlsx")
async def export_event_excel(
    _: Principal = Depends(require_admin),
    export_session: AsyncSession = Depends(get_export_session),
):
    """Download a human-readable operational backup sourced only from PostgreSQL."""
    exported_at = datetime.now(timezone.utc)
    rounds = (await export_session.execute(select(RoundState).order_by(RoundState.id))).scalars().all()
    teams = (await export_session.execute(select(Team).order_by(Team.name, Team.id))).scalars().all()
    users = (await export_session.execute(select(User).order_by(User.username, User.id))).scalars().all()
    questions = (await export_session.execute(select(Question).order_by(Question.index, Question.id))).scalars().all()
    team_questions = (await export_session.execute(select(TeamQuestion).order_by(TeamQuestion.question_id, TeamQuestion.team_id))).scalars().all()
    audit_events = (await export_session.execute(select(AuditEvent).order_by(AuditEvent.at, AuditEvent.id))).scalars().all()
    auth_sessions = (await export_session.execute(select(AuthSession).order_by(AuthSession.created_at, AuthSession.id))).scalars().all()
    swaps = (await export_session.execute(select(SwapEvent).order_by(SwapEvent.created_at, SwapEvent.id))).scalars().all()
    published = await export_session.get(PublishedLeaderboard, 1)

    round_state = next((row for row in rounds if row.id == 1), None)
    teams_by_id = {str(team.id): team for team in teams}
    users_by_id = {str(user.id): user for user in users}
    questions_by_id = {str(question.id): question for question in questions}
    current_question = questions_by_id.get(str(round_state.current_question_id)) if round_state and round_state.current_question_id else None

    workbook = Workbook()
    workbook.remove(workbook.active)
    overview_rows = [
        ["export_timestamp_utc", exported_at], ["application", "levelup"],
        ["round_phase", round_state.phase if round_state else None], ["round_version", round_state.version if round_state else None],
        ["current_question_id", round_state.current_question_id if round_state else None], ["current_question_index", current_question.index if current_question else None],
        ["live_started_at", round_state.live_started_at if round_state else None], ["global_deadline_at", round_state.global_deadline_at if round_state else None],
        ["paused_at", round_state.paused_at if round_state else None], ["team_count", len(teams)], ["user_count", len(users)],
        ["team_question_count", len(team_questions)], ["audit_event_count", len(audit_events)],
        ["published_at", published.published_at if published else None], ["published_phase", published.phase if published else None],
        ["published_question_index", published.question_index if published else None], ["published_round_version", published.round_version if published else None],
    ]
    overview = _add_excel_sheet(workbook, "Overview", ["Field", "Value"], overview_rows)
    overview.auto_filter.ref = None
    _add_excel_sheet(workbook, "Teams", ["team_id", "team_name", "current_score", "swap_available", "version", "created_at"], [
        [team.id, team.name, team.score, team.swap_available, team.version, team.created_at] for team in teams
    ])
    _add_excel_sheet(workbook, "Team Questions", ["team_id", "team_name", "question_id", "question_index", "answer", "is_correct", "points_delta", "graded", "locked", "locked_elapsed_ms", "freeze_elapsed_ms", "deadline_at", "bounty", "freeze", "peek", "swap_used", "kingslayer_eligible", "admin_override"], [
        [row.team_id, teams_by_id.get(str(row.team_id)).name if str(row.team_id) in teams_by_id else None, row.question_id, questions_by_id.get(str(row.question_id)).index if str(row.question_id) in questions_by_id else None, row.answer, row.is_correct, row.points_delta, row.graded, row.locked, row.locked_elapsed_ms, row.freeze_elapsed_ms, row.deadline_at, row.bounty, row.freeze, row.peek, row.swap_used, row.kingslayer_eligible, row.admin_override] for row in team_questions
    ])
    _add_excel_sheet(workbook, "Audit Log", ["timestamp", "event_type", "actor", "team_id", "team_name", "question_id", "question_index", "metadata"], [
        [row.at, row.kind, row.actor, row.team_id, teams_by_id.get(str(row.team_id)).name if row.team_id and str(row.team_id) in teams_by_id else None, row.question_id, questions_by_id.get(str(row.question_id)).index if row.question_id and str(row.question_id) in questions_by_id else None, row.payload] for row in audit_events
    ])
    _add_excel_sheet(workbook, "Sessions", ["session_id", "team_id", "team_name", "user_id", "username", "slot", "created_at", "last_seen_at", "expires_at", "revoked_at", "status"], [
        [row.id, row.team_id, teams_by_id.get(str(row.team_id)).name if row.team_id and str(row.team_id) in teams_by_id else None, row.user_id, users_by_id.get(str(row.user_id)).username if str(row.user_id) in users_by_id else None, row.slot, row.created_at, row.last_seen_at, row.expires_at, row.revoked_at, _session_export_status(row, exported_at)] for row in auth_sessions
    ])
    _add_excel_sheet(workbook, "Published Leaderboard", ["rank", "team_id", "team_name", "score"], [
        [row.get("rank"), row.get("team_id", row.get("id")), row.get("team_name", row.get("name")), row.get("score")] for row in (published.rows if published else [])
    ])
    _add_excel_sheet(workbook, "Card Usage", ["team_id", "team_name", "bounty_used", "freeze_used", "peek_used", "swap_used", "swap_target_team_id", "swap_target_team_name"], [
        [row.get("team_id"), row.get("team_name"), row.get("bounty_used"), row.get("freeze_used"), row.get("peek_used"), row.get("swap_used"), row.get("swap_target_team_id"), row.get("swap_target_team_name")] for row in (published.card_usage_rows if published else [])
    ])
    _add_excel_sheet(workbook, "Swap Events", ["created_at", "question_id", "question_index", "source_team_id", "source_team_name", "target_team_id", "target_team_name", "applied"], [
        [row.created_at, row.question_id, questions_by_id.get(str(row.question_id)).index if str(row.question_id) in questions_by_id else None, row.source_team_id, teams_by_id.get(str(row.source_team_id)).name if str(row.source_team_id) in teams_by_id else None, row.target_team_id, teams_by_id.get(str(row.target_team_id)).name if str(row.target_team_id) in teams_by_id else None, row.applied] for row in swaps
    ])
    output = BytesIO()
    workbook.save(output)
    filename = f"delta_x_round2_backup_{exported_at.strftime('%Y-%m-%d_%H%M')}_UTC.xlsx"
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.post("/round/reveal-cards/{on}")
async def reveal_cards(
    on: int,
    session: AsyncSession = Depends(get_session),
    _: Principal = Depends(require_admin),
):
    rnd = await get_round(session, for_update=True)
    rnd.reveal_cards = bool(on)
    await session.commit()
    return {"reveal_cards": rnd.reveal_cards}


def _public_solution(question: Question | None) -> dict[str, Any] | None:
    """Whitelist only audience-safe MCQ fields for durable publication."""
    if question is None or question.qtype not in {"mcq_single", "mcq_multi"}:
        return None
    answer_key = question.answer_key if isinstance(question.answer_key, dict) else {}
    if question.qtype == "mcq_single":
        raw_ids = [answer_key.get("option_id")]
    else:
        value = answer_key.get("option_ids")
        raw_ids = value if isinstance(value, list) else []
    return {
        "question_index": question.index,
        "type": question.qtype,
        "correct_option_ids": [str(value) for value in raw_ids if value is not None],
        "hint": question.hint if question.hint else None,
    }


@router.post("/round/publish-leaderboard")
async def api_publish_leaderboard(
    session: AsyncSession = Depends(get_session),
    admin: Principal = Depends(require_admin),
):
    from app.redisutil import set_frozen_leaderboard

    rnd = await get_round(session, for_update=True)
    question = await session.get(Question, rnd.current_question_id) if rnd.current_question_id else None
    published_solution = _public_solution(question)
    live_board = await leaderboard(session)
    card_usage_rows = []
    if rnd.current_question_id is not None:
        question_rows = (
            await session.execute(
                select(TeamQuestion).where(
                    TeamQuestion.question_id == rnd.current_question_id
                )
            )
        ).scalars().all()
        cards_by_team = {str(row.team_id): row for row in question_rows}

        swap_events = (
            await session.execute(
                select(SwapEvent).where(
                    SwapEvent.question_id == rnd.current_question_id
                )
            )
        ).scalars().all()

        swap_target_by_source = {
            str(event.source_team_id): str(event.target_team_id)
            for event in swap_events
        }

        team_name_by_id = {
            team["id"]: team["name"]
            for team in live_board
        }

        card_usage_rows = [
            {
                "team_id": team["id"],
                "team_name": team["name"],
                "bounty_used": bool(cards_by_team.get(team["id"]) and cards_by_team[team["id"]].bounty),
                "freeze_used": bool(cards_by_team.get(team["id"]) and cards_by_team[team["id"]].freeze),
                "peek_used": bool(cards_by_team.get(team["id"]) and cards_by_team[team["id"]].peek),
                "swap_used": bool(cards_by_team.get(team["id"]) and cards_by_team[team["id"]].swap_used),
                "swap_target_team_id": swap_target_by_source.get(team["id"]),
                "swap_target_team_name": team_name_by_id.get(swap_target_by_source.get(team["id"])),
            }
            for team in live_board
        ]
    published_at = datetime.now(timezone.utc)
    snapshot = await session.get(PublishedLeaderboard, 1, with_for_update=True)
    if snapshot is None:
        snapshot = PublishedLeaderboard(id=1)
        session.add(snapshot)
    snapshot.published_at = published_at
    snapshot.published_by = admin.username
    snapshot.round_version = rnd.version
    snapshot.phase = rnd.phase
    snapshot.question_id = rnd.current_question_id
    snapshot.question_index = question.index if question else None
    snapshot.rows = live_board
    snapshot.card_usage_rows = card_usage_rows
    snapshot.published_solution = published_solution
    await audit(
        session,
        actor=admin.username,
        kind="leaderboard_publish",
        question_id=rnd.current_question_id,
        payload={
            "published_at": published_at.isoformat(),
            "round_version": rnd.version,
            "phase": rnd.phase,
            "question_id": str(rnd.current_question_id) if rnd.current_question_id else None,
            "question_index": question.index if question else None,
            "rows": live_board,
            "card_usage_rows": card_usage_rows,
            "published_solution": published_solution,
            "actor": admin.username,
        },
    )
    await session.commit()

    context = {"round_version": rnd.version, "published_at": published_at.isoformat()}
    await notify_after_commit(
        set_frozen_leaderboard(live_board),
        operation="admin.publish_leaderboard.cache",
        context=context,
    )
    await notify_after_commit(
        hub.broadcast(PUBLIC, {"type": "leaderboard", "leaderboard": live_board}),
        operation="admin.publish_leaderboard.broadcast",
        context=context,
    )

    return {"ok": True, "published": True}

# Ensure the upload directory exists
UPLOAD_DIR = "app/static/uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

@router.post("/upload")
async def upload_image(
    file: UploadFile = File(...),
    _: Principal = Depends(require_admin)
):
    # Save the file locally
    file_location = f"{UPLOAD_DIR}/{file.filename}"
    with open(file_location, "wb+") as file_object:
        shutil.copyfileobj(file.file, file_object)
    
    # Return the relative URL path to be saved in the database
    return {"media_url": f"/static/uploads/{file.filename}"}
